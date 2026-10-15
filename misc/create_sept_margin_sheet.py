"""
Generate Under-Margin Discrepancies sheet for September 2026
Matching the exact format of the August sheet shown in the screenshot.

Columns:
A: Invoice #
B: Stock Item Description
C: MRP (Rs)
D: Landed Cost (5% GST) = Purchase Rate * 1.05
E: Net Margin % = (MRP - Landed Cost) / MRP * 100
F: Expected Buy Rate @ 37.5% Disc. (Incl. GST) = MRP * 0.625
G: Difference = Expected Buy Rate - Landed Cost
H: Pairs
I: Total Difference = Difference * Pairs

Filter: Items where Net Margin % < 37.5% (i.e. purchase disc < ~40.48%)
"""
import json, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers

with open(r'c:\Projects\SKIL Master\sync-engine\data\vouchers.json', 'r', encoding='utf-8') as f:
    v = json.load(f)

purchases = {int(x['voucherNumber']): x for x in v.values() if x.get('voucherType') == 'Purchase' and str(x.get('voucherNumber')).isdigit()}

def clean_rate(rate_str):
    if not rate_str:
        return 0.0
    r = rate_str.split('/')[0].strip().replace(',', '')
    try:
        return float(r)
    except:
        return 0.0

def get_mrp(item_name):
    m = re.search(r'MRP\s*(\d+(\.\d+)?)/-', item_name, re.IGNORECASE)
    if m:
        return float(m.group(1))
    return None

def get_pairs(qty_str):
    m = re.search(r'(\d+)\s*Pairs?', str(qty_str), re.IGNORECASE)
    if m:
        return int(m.group(1))
    return 0

# September vouchers: P#23 to P#53
sept_range = range(23, 54)

# Collect under-margin items
under_margin_items = []

for i in sept_range:
    p = purchases.get(i)
    if not p:
        continue
    pref = p.get('reference', '').strip()
    
    # Skip non-Ajanta invoices (XPL brand has different pricing)
    if pref.startswith('XPL'):
        continue
    
    for it in p.get('inventoryEntries', []):
        iname = it.get('stockItemName', '')
        prate = clean_rate(it.get('rate', ''))
        mrp = get_mrp(iname)
        pairs = get_pairs(it.get('billedQty', ''))
        
        if not mrp or mrp <= 0 or prate <= 0:
            continue
        
        landed_cost = round(prate * 1.05, 2)  # Purchase Rate + 5% IGST
        net_margin_pct = round(((mrp - landed_cost) / mrp) * 100, 2)
        expected_buy_rate = round(mrp * 0.625, 2)  # 37.5% discount on MRP
        difference = round(expected_buy_rate - landed_cost, 2)
        total_difference = round(difference * pairs, 2)
        
        # Filter: Net Margin < 37.5% AND difference is meaningful (>= Rs 0.50 per pair)
        # This excludes 37.49% margin items that are just paise-level rounding
        if net_margin_pct < 37.50 and abs(difference) >= 0.50:
            # Build description similar to screenshot format
            # E.g.: "PUL 979E/CH+MR SZ-4X7-(30PR)"
            # Extract article code from item name: "AJ PUL 979E/CH+MR 04/07 MRP229.75/-"
            # -> "PUL 979E/CH+MR"
            desc_match = re.match(r'AJ\s+(.*?)\s+\d+/\d+\s+MRP', iname)
            if desc_match:
                article = desc_match.group(1)
            else:
                desc_match2 = re.match(r'AJ\s+(.*?)\s+MRP', iname)
                article = desc_match2.group(1) if desc_match2 else iname
            
            # Extract size info from item name: "04/07" -> "SZ-4X7"
            sz_match = re.search(r'(\d+)/(\d+)\s+MRP', iname)
            if sz_match:
                sz1 = sz_match.group(1).lstrip('0') or '0'
                sz2 = sz_match.group(2).lstrip('0') or '0'
                size_str = f"SZ-{sz1}X{sz2}"
            else:
                sz_match2 = re.search(r'(\d+)\s+MRP', iname)
                size_str = f"SZ-{sz_match2.group(1)}" if sz_match2 else ""
            
            desc = f"{article} {size_str}-({pairs} PR)" if size_str else f"{article} ({pairs} PR)"
            
            under_margin_items.append({
                'invoice': pref,
                'description': desc,
                'mrp': mrp,
                'landed_cost': landed_cost,
                'net_margin_pct': net_margin_pct,
                'expected_buy_rate': expected_buy_rate,
                'difference': difference,
                'pairs': pairs,
                'total_difference': total_difference,
                'vch_no': i
            })

# Sort by invoice number
under_margin_items.sort(key=lambda x: (x['invoice'], x['description']))

print(f"Total under-margin items found for September: {len(under_margin_items)}")
for item in under_margin_items:
    print(f"  {item['invoice']:>14} | {item['description'][:45]:45} | MRP: {item['mrp']:>7.2f} | Landed: {item['landed_cost']:>7.2f} | Margin: {item['net_margin_pct']:>6.2f}% | Diff: {item['difference']:>+8.2f} | Pairs: {item['pairs']:>3} | TotalDiff: {item['total_difference']:>+10.2f}")

# Create Excel workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Under Margin Discrepancies"

# Styles
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

data_font = Font(name='Arial', size=10)
red_font = Font(name='Arial', size=10, color='FF0000')
green_font = Font(name='Arial', size=10, color='008000')
red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='B0B0B0'),
    right=Side(style='thin', color='B0B0B0'),
    top=Side(style='thin', color='B0B0B0'),
    bottom=Side(style='thin', color='B0B0B0')
)

# Headers (Row 1)
headers = [
    "Invoice #",
    "Stock Item Description",
    "MRP\n(Rs)",
    "Landed Cost\n(5% GST)",
    "Net Margin %",
    "Expected Buy Rate\n@ 37.5% Disc. (Incl. GST)",
    "Difference",
    "Pairs",
    "Total\nDifference"
]

ws.row_dimensions[1].height = 36
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Data rows
for row_idx, item in enumerate(under_margin_items, 2):
    ws.row_dimensions[row_idx].height = 18
    
    # A: Invoice #
    c1 = ws.cell(row=row_idx, column=1, value=item['invoice'])
    c1.font = data_font
    c1.alignment = Alignment(horizontal='left', vertical='center')
    c1.border = thin_border
    
    # B: Stock Item Description
    c2 = ws.cell(row=row_idx, column=2, value=item['description'])
    c2.font = data_font
    c2.alignment = Alignment(horizontal='left', vertical='center')
    c2.border = thin_border
    
    # C: MRP
    c3 = ws.cell(row=row_idx, column=3, value=item['mrp'])
    c3.font = data_font
    c3.alignment = Alignment(horizontal='right', vertical='center')
    c3.number_format = '#,##0.00'
    c3.border = thin_border
    
    # D: Landed Cost
    c4 = ws.cell(row=row_idx, column=4, value=item['landed_cost'])
    c4.font = data_font
    c4.alignment = Alignment(horizontal='right', vertical='center')
    c4.number_format = '#,##0.00'
    c4.border = thin_border
    
    # E: Net Margin %
    c5 = ws.cell(row=row_idx, column=5, value=f"{item['net_margin_pct']:.2f}%")
    c5.alignment = Alignment(horizontal='right', vertical='center')
    c5.border = thin_border
    # Color: red if < 35%, orange-ish if < 37%, normal otherwise
    if item['net_margin_pct'] < 35.0:
        c5.font = Font(name='Arial', size=10, bold=True, color='FF0000')
        c5.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    elif item['net_margin_pct'] < 37.0:
        c5.font = Font(name='Arial', size=10, bold=True, color='CC6600')
    else:
        c5.font = Font(name='Arial', size=10, bold=True)
    
    # F: Expected Buy Rate
    c6 = ws.cell(row=row_idx, column=6, value=item['expected_buy_rate'])
    c6.font = data_font
    c6.alignment = Alignment(horizontal='right', vertical='center')
    c6.number_format = '#,##0.00'
    c6.border = thin_border
    
    # G: Difference
    c7 = ws.cell(row=row_idx, column=7, value=item['difference'])
    c7.font = red_font
    c7.alignment = Alignment(horizontal='right', vertical='center')
    c7.number_format = '#,##0.00'
    c7.border = thin_border
    
    # H: Pairs
    c8 = ws.cell(row=row_idx, column=8, value=item['pairs'])
    c8.font = data_font
    c8.alignment = Alignment(horizontal='center', vertical='center')
    c8.number_format = '#,##0'
    c8.border = thin_border
    
    # I: Total Difference
    c9 = ws.cell(row=row_idx, column=9, value=item['total_difference'])
    c9.font = Font(name='Arial', size=10, bold=True, color='FF0000')
    c9.alignment = Alignment(horizontal='right', vertical='center')
    c9.number_format = '#,##0.00'
    c9.border = thin_border

# Total row at the bottom
tot_row = len(under_margin_items) + 2
ws.row_dimensions[tot_row].height = 22

tot_font = Font(name='Arial', size=11, bold=True, color='FF0000')
tot_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

# Sum formula for Total Difference
ws.cell(row=tot_row, column=9, value=f"=SUM(I2:I{tot_row-1})")
ws.cell(row=tot_row, column=9).font = tot_font
ws.cell(row=tot_row, column=9).fill = tot_fill
ws.cell(row=tot_row, column=9).number_format = '#,##0.00'
ws.cell(row=tot_row, column=9).alignment = Alignment(horizontal='right', vertical='center')
ws.cell(row=tot_row, column=9).border = thin_border

# Column widths
col_widths = {'A': 16, 'B': 42, 'C': 10, 'D': 14, 'E': 14, 'F': 28, 'G': 12, 'H': 8, 'I': 14}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# Auto-filter
ws.auto_filter.ref = f"A1:I{tot_row-1}"

# Save
output_path = r'c:\Projects\SKIL Master\SKIL_Lifestyle_Under_Margin_Discrepancies_2026-09-01_to_2026-09-30.xlsx'
output_downloads = r'C:\Users\Sahil Kumar\Downloads\SKIL_Lifestyle_Under_Margin_Discrepancies_2026-09-01_to_2026-09-30.xlsx'

wb.save(output_path)
wb.save(output_downloads)
print(f"\nSaved to:\n  1. {output_path}\n  2. {output_downloads}")

# Print grand total
grand_total = sum(item['total_difference'] for item in under_margin_items)
print(f"\nGrand Total Difference: Rs {grand_total:,.2f}")

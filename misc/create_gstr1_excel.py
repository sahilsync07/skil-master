import json, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

with open(r'c:\Projects\SKIL Master\sync-engine\data\vouchers.json', 'r', encoding='utf-8') as f:
    v = json.load(f)

sales = [x for x in v.values() if x.get('voucherType') == 'Sales' and '20260901' <= x.get('date', '') <= '20260930']
sales.sort(key=lambda x: int(str(x.get('voucherNumber')).replace('SKIL-', '')))

gstin_map = {
    'SRI SHARVANI TRADERS': ('37ARWPY2407N1ZR', '37-Andhra Pradesh'),
    'K G N FOOTWARE': ('37CPZPK0089M1ZW', '37-Andhra Pradesh'),
    'M/S.SHREE FOOTWEAR': ('37AEBPP0086L1Z9', '37-Andhra Pradesh'),
    'JYOTHI FOOT WEAR': ('37AUQPP0952Q1ZO', '37-Andhra Pradesh'),
    'M/s Sri Brundabana Enterprises, Rayagada': ('21ADKPC5405A1ZI', '21-Odisha'),
    'M/S SRI BRUNDABAN ENTERPRISES, MALKANGIRI': ('21ADKPC5405A1ZI', '21-Odisha'),
    'MAHAVIR SHOE HOUSE': ('21BGYPB6336B1ZO', '21-Odisha'),
    'M/S SUSHILA BOOT HOUSE': ('21AZKPP2729C1ZN', '21-Odisha')
}

wb = openpyxl.Workbook()

# Styles
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
data_font = Font(name='Segoe UI', size=10)
bold_font = Font(name='Segoe UI', size=10, bold=True)
tot_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# ----------------- TAB 1: B2B Invoices (Table 4A) -----------------
ws1 = wb.active
ws1.title = "B2B Invoices"

headers1 = [
    "GSTIN/UIN of Recipient", "Receiver Name", "Invoice Number", "Invoice date",
    "Invoice Value", "Place Of Supply", "Reverse Charge", "Applicable % of Tax Rate",
    "Invoice Type", "E-Commerce GSTIN", "Rate (%)", "Taxable Value", "Cess Amount"
]

ws1.row_dimensions[1].height = 28
for col, h in enumerate(headers1, 1):
    c = ws1.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

row = 2
for s in sales:
    vno = s.get('voucherNumber')
    raw_date = s.get('date')
    d_fmt = f"{raw_date[6:8]}-{raw_date[4:6]}-{raw_date[:4]}"
    tot_amt = abs(float(s.get('amount', 0)))
    
    party = s.get('partyLedgerName')
    if not party:
        for le in s.get('ledgerEntries', []):
            ln = le.get('ledgerName', '')
            if ln in gstin_map:
                party = ln
                break
    
    gstin, pos = gstin_map.get(party, ('', ''))
    taxable = sum(abs(float(le.get('amount', 0))) for le in s.get('ledgerEntries', []) if 'Sales' in le.get('ledgerName', ''))
    
    ws1.cell(row=row, column=1, value=gstin).font = data_font
    ws1.cell(row=row, column=2, value=party).font = data_font
    ws1.cell(row=row, column=3, value=vno).font = data_font
    ws1.cell(row=row, column=4, value=d_fmt).font = data_font
    
    c5 = ws1.cell(row=row, column=5, value=tot_amt)
    c5.font = data_font
    c5.number_format = '#,##0.00'
    
    ws1.cell(row=row, column=6, value=pos).font = data_font
    ws1.cell(row=row, column=7, value="N").font = data_font
    ws1.cell(row=row, column=8, value="").font = data_font
    ws1.cell(row=row, column=9, value="Regular").font = data_font
    ws1.cell(row=row, column=10, value="").font = data_font
    ws1.cell(row=row, column=11, value=5).font = data_font
    
    c12 = ws1.cell(row=row, column=12, value=taxable)
    c12.font = data_font
    c12.number_format = '#,##0.00'
    
    ws1.cell(row=row, column=13, value=0.0).font = data_font
    
    for c_idx in range(1, 14):
        ws1.cell(row=row, column=c_idx).border = thin_border
    row += 1

# Total row
ws1.cell(row=row, column=2, value="TOTAL").font = bold_font
ws1.cell(row=row, column=2).fill = tot_fill

c_tot_inv = ws1.cell(row=row, column=5, value=f"=SUM(E2:E{row-1})")
c_tot_inv.font = bold_font
c_tot_inv.fill = tot_fill
c_tot_inv.number_format = '#,##0.00'

c_tot_tax = ws1.cell(row=row, column=12, value=f"=SUM(L2:L{row-1})")
c_tot_tax.font = bold_font
c_tot_tax.fill = tot_fill
c_tot_tax.number_format = '#,##0.00'

for c_idx in range(1, 14):
    ws1.cell(row=row, column=c_idx).border = thin_border

# Auto column widths
for col in ws1.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

# ----------------- TAB 2: Party Summary -----------------
ws2 = wb.create_sheet(title="Party Summary")
headers2 = ["GSTIN", "Receiver Name", "State / POS", "Bills", "Taxable Value", "Rate", "IGST", "CGST", "SGST", "Total Invoice Value"]

ws2.row_dimensions[1].height = 28
for col, h in enumerate(headers2, 1):
    c = ws2.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')

from collections import defaultdict
p_grp = defaultdict(lambda: {'count': 0, 'taxable': 0.0, 'igst': 0.0, 'cgst': 0.0, 'sgst': 0.0, 'total': 0.0})

for s in sales:
    party = s.get('partyLedgerName')
    if not party:
        for le in s.get('ledgerEntries', []):
            ln = le.get('ledgerName', '')
            if ln in gstin_map:
                party = ln; break
    tot = abs(float(s.get('amount', 0)))
    taxable = sum(abs(float(le.get('amount', 0))) for le in s.get('ledgerEntries', []) if 'Sales' in le.get('ledgerName', ''))
    igst = sum(abs(float(le.get('amount', 0))) for le in s.get('ledgerEntries', []) if 'Output IGST' in le.get('ledgerName', '') or ('IGST' in le.get('ledgerName', '') and 'Input' not in le.get('ledgerName', '')))
    cgst = sum(abs(float(le.get('amount', 0))) for le in s.get('ledgerEntries', []) if 'Output CGST' in le.get('ledgerName', '') or ('CGST' in le.get('ledgerName', '') and 'Input' not in le.get('ledgerName', '')))
    sgst = sum(abs(float(le.get('amount', 0))) for le in s.get('ledgerEntries', []) if 'Output SGST' in le.get('ledgerName', '') or ('SGST' in le.get('ledgerName', '') and 'Input' not in le.get('ledgerName', '')))
    p_grp[party]['count'] += 1
    p_grp[party]['taxable'] += taxable
    p_grp[party]['igst'] += igst
    p_grp[party]['cgst'] += cgst
    p_grp[party]['sgst'] += sgst
    p_grp[party]['total'] += tot

r = 2
for party, d in sorted(p_grp.items()):
    gstin, pos = gstin_map.get(party, ('', ''))
    ws2.cell(row=r, column=1, value=gstin).font = data_font
    ws2.cell(row=r, column=2, value=party).font = data_font
    ws2.cell(row=r, column=3, value=pos).font = data_font
    ws2.cell(row=r, column=4, value=d['count']).font = data_font
    
    ws2.cell(row=r, column=5, value=d['taxable']).number_format = '#,##0.00'
    ws2.cell(row=r, column=6, value="5%").font = data_font
    ws2.cell(row=r, column=7, value=d['igst']).number_format = '#,##0.00'
    ws2.cell(row=r, column=8, value=d['cgst']).number_format = '#,##0.00'
    ws2.cell(row=r, column=9, value=d['sgst']).number_format = '#,##0.00'
    ws2.cell(row=r, column=10, value=d['total']).number_format = '#,##0.00'
    
    for c_idx in range(1, 11):
        ws2.cell(row=r, column=c_idx).border = thin_border
    r += 1

# Total row
ws2.cell(row=r, column=2, value="TOTAL").font = bold_font
ws2.cell(row=r, column=2).fill = tot_fill
ws2.cell(row=r, column=4, value=f"=SUM(D2:D{r-1})").font = bold_font
ws2.cell(row=r, column=4).fill = tot_fill
for c_idx, col_l in [(5, 'E'), (7, 'G'), (8, 'H'), (9, 'I'), (10, 'J')]:
    c = ws2.cell(row=r, column=c_idx, value=f"=SUM({col_l}2:{col_l}{r-1})")
    c.font = bold_font
    c.fill = tot_fill
    c.number_format = '#,##0.00'
for c_idx in range(1, 11):
    ws2.cell(row=r, column=c_idx).border = thin_border

for col in ws2.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

# ----------------- TAB 3: HSN & Document Summary -----------------
ws3 = wb.create_sheet(title="HSN & Docs Summary")

# Section 1: Table 12 HSN Summary
ws3.cell(row=1, column=1, value="TABLE 12: HSN-WISE SUMMARY OF OUTWARD SUPPLIES").font = Font(name='Segoe UI', size=12, bold=True, color='1F4E79')
headers_hsn = ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Taxable Value", "Integrated Tax", "Central Tax", "State Tax", "Cess Amount"]
for col, h in enumerate(headers_hsn, 1):
    c = ws3.cell(row=2, column=col, value=h)
    c.font = header_font; c.fill = header_fill; c.alignment = Alignment(horizontal='center', vertical='center')

total_pairs = 10918
ws3.cell(row=3, column=1, value="6402").font = data_font
ws3.cell(row=3, column=2, value="Footwear with outer soles and uppers of rubber or plastics").font = data_font
ws3.cell(row=3, column=3, value="PRS-PAIRS").font = data_font
ws3.cell(row=3, column=4, value=total_pairs).font = data_font
ws3.cell(row=3, column=5, value=1488717.00).number_format = '#,##0.00'
ws3.cell(row=3, column=6, value=1417825.82).number_format = '#,##0.00'
ws3.cell(row=3, column=7, value=30603.18).number_format = '#,##0.00'
ws3.cell(row=3, column=8, value=20144.02).number_format = '#,##0.00'
ws3.cell(row=3, column=9, value=20144.02).number_format = '#,##0.00'
ws3.cell(row=3, column=10, value=0.0).number_format = '#,##0.00'

for c_idx in range(1, 11):
    ws3.cell(row=3, column=c_idx).border = thin_border

# Section 2: Table 13 Documents Summary
ws3.cell(row=6, column=1, value="TABLE 13: DOCUMENTS ISSUED DURING THE TAX PERIOD").font = Font(name='Segoe UI', size=12, bold=True, color='1F4E79')
headers_doc = ["Nature of Document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled", "Net Issued"]
for col, h in enumerate(headers_doc, 1):
    c = ws3.cell(row=7, column=col, value=h)
    c.font = header_font; c.fill = header_fill; c.alignment = Alignment(horizontal='center', vertical='center')

ws3.cell(row=8, column=1, value="Invoices for outward supply").font = data_font
ws3.cell(row=8, column=2, value="SKIL-023").font = data_font
ws3.cell(row=8, column=3, value="SKIL-053").font = data_font
ws3.cell(row=8, column=4, value=31).font = data_font
ws3.cell(row=8, column=5, value=0).font = data_font
ws3.cell(row=8, column=6, value=31).font = data_font

for c_idx in range(1, 7):
    ws3.cell(row=8, column=c_idx).border = thin_border

for col in ws3.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)

file1 = r'c:\Projects\SKIL Master\SKIL_GSTR1_September_2026.xlsx'
file2 = r'C:\Users\Sahil Kumar\Downloads\SKIL_GSTR1_September_2026.xlsx'
wb.save(file1)
wb.save(file2)

print(f"GSTR-1 Excel files successfully saved to:")
print(f"  1. {file1}")
print(f"  2. {file2}")

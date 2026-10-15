import fitz, glob, os, json, re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 1. Load data
purchases_dir = glob.glob(r'c:\Projects\SKIL Master\user-documents-read-only-access\Purchases*')[0]

with open(r'c:\Projects\SKIL Master\sync-engine\data\vouchers.json', 'r', encoding='utf-8') as f:
    v = json.load(f)

purchases = {int(x['voucherNumber']): x for x in v.values() if x.get('voucherType') == 'Purchase' and str(x.get('voucherNumber')).isdigit()}
sales = {}
for x in v.values():
    if x.get('voucherType') == 'Sales':
        vn = str(x.get('voucherNumber')).replace('SKIL-', '')
        if vn.isdigit():
            sales[int(vn)] = x

pdf_files = [f for f in os.listdir(purchases_dir) if f.endswith('.pdf')]

def parse_pdf(filepath):
    doc = fitz.open(filepath)
    txt = "\n".join([page.get_text() for page in doc])
    
    m_inv = re.search(r'Inv\.\s*No\.\s*:\s*([^\n\r]+)', txt)
    inv_no = m_inv.group(1).strip() if m_inv else None
    
    m_date = re.search(r'Inv\.\s*Date\s*:\s*([^\n\r]+)', txt)
    inv_date = m_date.group(1).strip() if m_date else None
    
    m_ctn = re.search(r'To\s*tal\s*:\s*\n\s*(\d+)\s*\n\s*(\d+)', txt, re.IGNORECASE)
    ctns = int(m_ctn.group(1)) if m_ctn else None
    
    if not ctns:
        m_ctn2 = re.search(r'Total\s*Carton[s]?\s*[:\s]*(\d+)', txt, re.IGNORECASE)
        if m_ctn2:
            ctns = int(m_ctn2.group(1))
            
    ship_to = "UNKNOWN"
    m_ship = re.search(r'Ship-To:\s*\n(.*?)\n(?:Bill-To|Order No)', txt, re.DOTALL | re.IGNORECASE)
    if m_ship:
        s_block = m_ship.group(1)
        s_low = s_block.lower()
        if 'malkangiri' in s_low or 'malakangiri' in s_low:
            ship_to = 'Sri Brundabana Enterprises, Malkangiri'
        elif 'brundaban' in s_low:
            ship_to = 'Sri Brundabana Enterprises, Rayagada'
        elif 'shree footwear' in s_low:
            ship_to = 'M/s. Shree Footwear, Srikakulam'
        elif 'delight' in s_low:
            ship_to = 'Delight Footwear, Vizianagaram'
        elif 'k g n' in s_low or 'kgn' in s_low:
            ship_to = 'K G N Footware, Parvathipuram'
        elif 'sharvani' in s_low:
            ship_to = 'Sri Sharvani Traders, Visakhapatnam'
        elif 'jyothi' in s_low:
            ship_to = 'Jyothi Footwear, Vizianagaram'
        elif 'sushila' in s_low:
            ship_to = 'M/s Sushila Boot House, Jharsuguda'
        elif 'mahavir' in s_low:
            ship_to = 'Mahavir Shoe House, Dhenkanal'
            
    return {
        'inv_no': inv_no,
        'inv_date': inv_date,
        'ctns': ctns,
        'ship_to': ship_to
    }

pdf_data = {}
for pf in pdf_files:
    info = parse_pdf(os.path.join(purchases_dir, pf))
    if info['inv_no']:
        pdf_data[info['inv_no']] = info
        pdf_data[info['inv_no'].replace('/', '-')] = info

def get_consignee(code_or_party):
    cp = code_or_party.upper()
    if 'MALKANGIRI' in cp or 'SBEM' in cp:
        return 'Sri Brundabana Enterprises, Malkangiri'
    elif 'BRUNDABAN' in cp or 'SBE' in cp:
        return 'Sri Brundabana Enterprises, Rayagada'
    elif 'SHREE' in cp or 'SF' in cp:
        return 'M/s. Shree Footwear, Srikakulam'
    elif 'DELIGHT' in cp:
        return 'Delight Footwear, Vizianagaram'
    elif 'K G N' in cp or 'KGN' in cp:
        return 'K G N Footware, Parvathipuram'
    elif 'SHARVANI' in cp:
        return 'Sri Sharvani Traders, Visakhapatnam'
    elif 'JYOTHI' in cp:
        return 'Jyothi Footwear, Vizianagaram'
    elif 'SUSHILA' in cp or 'SBH' in cp:
        return 'M/s Sushila Boot House, Jharsuguda'
    elif 'MAHAVIR' in cp:
        return 'Mahavir Shoe House, Dhenkanal'
    return code_or_party

# Customer 5% CD allocation groups (Total = Rs 59,876.00)
cn_given_targets = {
    'DELIGHT': {'cn_amt': 7436.00, 'vnos': [2]},
    'KGN': {'cn_amt': 5421.00, 'vnos': [4, 5, 6]},
    'SF': {'cn_amt': 26539.00, 'vnos': list(range(11, 21))},
    'SBE': {'cn_amt': 20480.00, 'vnos': [1, 3, 7, 8, 9, 10, 21, 22]}
}

cn_given_map = {}
for gname, ginfo in cn_given_targets.items():
    vnos = ginfo['vnos']
    tot_cn = ginfo['cn_amt']
    tax_map = {}
    sum_tax = 0.0
    for vno in vnos:
        s = sales.get(vno)
        stax = sum([abs(float(it.get('amount', 0))) for it in s.get('inventoryEntries', [])])
        tax_map[vno] = stax
        sum_tax += stax
    
    allocated = 0.0
    for idx_v, vno in enumerate(vnos):
        if idx_v == len(vnos) - 1:
            val = round(tot_cn - allocated, 2)
        else:
            val = round((tax_map[vno] / sum_tax) * tot_cn, 2)
            allocated += val
        cn_given_map[vno] = val

tot_aug_ptax = 1192130.34
tot_aug_turnover_cn = 25727.00 + 37552.12 + 45896.00 # Rs 109,175.12

cn_rec_map = {}
for i in range(1, 23):
    p = purchases.get(i)
    ptax = sum([abs(float(it.get('amount', 0))) for it in p.get('inventoryEntries', [])])
    share = round((ptax / tot_aug_ptax) * tot_aug_turnover_cn, 2)
    if p.get('reference') == '1900405893': # P#19
        share += 1008.54
    elif p.get('reference') == '1900402693': # P#05
        share += 30.19
    cn_rec_map[i] = round(share, 2)

# Build 53 records
records = []
for i in range(1, 54):
    p = purchases.get(i)
    s = sales.get(i)
    ref = p.get('reference', '').strip() if p else ''
    pdate_raw = str(p.get('date', '')) if p else ''
    if len(pdate_raw) == 8 and pdate_raw.isdigit():
        pdate_dt = datetime.strptime(pdate_raw, '%Y%m%d')
        pdate_str = pdate_dt.strftime('%d-%m-%Y')
    else:
        pdate_dt = datetime(2026, 1, 1)
        pdate_str = pdate_raw

    pnarr = p.get('narration', '') if p else ''
    snarr = s.get('narration', '') if s else ''
    cust = s.get('partyName', '') if s else ''
    p_amt = abs(float(p.get('amount', 0))) if p else 0.0
    s_amt = abs(float(s.get('amount', 0))) if s else 0.0
    
    pdf_info = pdf_data.get(ref) or pdf_data.get(ref.replace('/', '-'))
    
    ctns = None
    if pdf_info and pdf_info.get('ctns'):
        ctns = pdf_info['ctns']
    else:
        m_c = re.search(r'(\d+)\s*(?:ctn|ctns)', pnarr, re.IGNORECASE) or re.search(r'(\d+)\s*(?:ctn|ctns)', snarr, re.IGNORECASE)
        if m_c:
            ctns = int(m_c.group(1))
            
    if ref == '1900405893': # Voucher 19
        ctns = 11

    if pdf_info and pdf_info.get('ship_to') and pdf_info['ship_to'] != 'UNKNOWN':
        consignee = pdf_info['ship_to']
    else:
        consignee = get_consignee(cust or pnarr or snarr)

    cn_given = cn_given_map.get(i, 0.0)
    cn_rec = cn_rec_map.get(i, 0.0)

    records.append({
        'vch_idx': i,
        'date_dt': pdate_dt,
        'date_str': pdate_str,
        'inv_no': ref,
        'p_vno': p.get('voucherNumber', str(i)),
        's_vno': s.get('voucherNumber', f"SKIL-{i:03d}"),
        'consignee': consignee,
        'ctns': ctns if ctns else 0,
        'p_amt': p_amt,
        's_amt': s_amt,
        'cn_rec': cn_rec,
        'cn_given': cn_given
    })

wb = openpyxl.Workbook()

# Setup styling objects
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

font_regular = Font(name='Segoe UI', size=10)
fill_even = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

tot_font = Font(name='Segoe UI', size=11, bold=True, color='000000')
tot_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
double_bottom_border = Border(
    left=Side(style='thin', color='B0B0B0'),
    right=Side(style='thin', color='B0B0B0'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='double', color='000000')
)

headers = [
    "Serial No",
    "Date",
    "Invoice No",
    "Purchase Voucher No in Tally",
    "Sales Voucher No in Tally",
    "End Consignee Name",
    "No of Ctns",
    "Purchase Amount",
    "Sale Amount",
    "Credit Notes Received",
    "Credit Note Given"
]

col_widths = {
    'A': 12, 'B': 14, 'C': 18, 'D': 28, 'E': 26, 'F': 42,
    'G': 14, 'H': 20, 'I': 20, 'J': 24, 'K': 22
}

def populate_sheet(ws, title_text, subtitle_text, row_data, serial_numbers):
    ws.views.sheetView[0].showGridLines = True
    
    # Title
    ws.merge_cells('A1:K1')
    ws['A1'] = title_text
    ws['A1'].font = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40

    ws.merge_cells('A2:K2')
    ws['A2'] = subtitle_text
    ws['A2'].font = Font(name='Segoe UI', size=10, italic=True, color='FFFFFF')
    ws['A2'].fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 22

    # Headers
    ws.row_dimensions[4].height = 32
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    start_row = 5
    for idx, (r, sno) in enumerate(zip(row_data, serial_numbers), 1):
        cur_row = start_row + idx - 1
        ws.row_dimensions[cur_row].height = 20
        row_fill = fill_even if idx % 2 == 0 else fill_odd
        
        c1 = ws.cell(row=cur_row, column=1, value=sno)
        c1.alignment = Alignment(horizontal='center', vertical='center')
        
        c2 = ws.cell(row=cur_row, column=2, value=r['date_str'])
        c2.alignment = Alignment(horizontal='center', vertical='center')
        
        c3 = ws.cell(row=cur_row, column=3, value=r['inv_no'])
        c3.alignment = Alignment(horizontal='center', vertical='center')
        
        c4 = ws.cell(row=cur_row, column=4, value=r['p_vno'])
        c4.alignment = Alignment(horizontal='center', vertical='center')
        
        c5 = ws.cell(row=cur_row, column=5, value=r['s_vno'])
        c5.alignment = Alignment(horizontal='center', vertical='center')
        
        c6 = ws.cell(row=cur_row, column=6, value=r['consignee'])
        c6.alignment = Alignment(horizontal='left', vertical='center')
        
        c7 = ws.cell(row=cur_row, column=7, value=r['ctns'])
        c7.alignment = Alignment(horizontal='center', vertical='center')
        c7.number_format = '#,##0'
        
        c8 = ws.cell(row=cur_row, column=8, value=r['p_amt'])
        c8.alignment = Alignment(horizontal='right', vertical='center')
        c8.number_format = '₹#,##0.00'
        
        c9 = ws.cell(row=cur_row, column=9, value=r['s_amt'])
        c9.alignment = Alignment(horizontal='right', vertical='center')
        c9.number_format = '₹#,##0.00'
        
        c10 = ws.cell(row=cur_row, column=10, value=r['cn_rec'])
        c10.alignment = Alignment(horizontal='right', vertical='center')
        c10.number_format = '₹#,##0.00'
        
        c11 = ws.cell(row=cur_row, column=11, value=r['cn_given'])
        c11.alignment = Alignment(horizontal='right', vertical='center')
        c11.number_format = '₹#,##0.00'
        
        for c in [c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, c11]:
            c.font = font_regular
            c.fill = row_fill
            c.border = thin_border

    tot_row = start_row + len(row_data)
    ws.row_dimensions[tot_row].height = 26

    ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=6)
    ws.cell(row=tot_row, column=1, value="TOTAL (53 BILLS)")
    ws.cell(row=tot_row, column=1).alignment = Alignment(horizontal='center', vertical='center')

    ws.cell(row=tot_row, column=7, value=f"=SUM(G{start_row}:G{tot_row-1})").number_format = '#,##0'
    ws.cell(row=tot_row, column=8, value=f"=SUM(H{start_row}:H{tot_row-1})").number_format = '₹#,##0.00'
    ws.cell(row=tot_row, column=9, value=f"=SUM(I{start_row}:I{tot_row-1})").number_format = '₹#,##0.00'
    ws.cell(row=tot_row, column=10, value=f"=SUM(J{start_row}:J{tot_row-1})").number_format = '₹#,##0.00'
    ws.cell(row=tot_row, column=11, value=f"=SUM(K{start_row}:K{tot_row-1})").number_format = '₹#,##0.00'

    for col_idx in range(1, 12):
        c = ws.cell(row=tot_row, column=col_idx)
        c.font = tot_font
        c.fill = tot_fill
        c.border = double_bottom_border
        if col_idx >= 7:
            c.alignment = Alignment(horizontal='right' if col_idx != 7 else 'center', vertical='center')

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

# --- TAB 1: REVERSE ORDER (Most Recent First: Bill #53 down to #1, Serial No: 53 down to 1) ---
ws1 = wb.active
ws1.title = "Reverse Order (53 to 1)"
# Sort descending by voucher index: 53 down to 1
rec_desc = sorted(records, key=lambda x: x['vch_idx'], reverse=True)
snos_desc = [r['vch_idx'] for r in rec_desc] # 53, 52, 51... 1
populate_sheet(
    ws1,
    "SKIL LIFESTYLE LLP — ALL 53 BILLS AUDIT REGISTER",
    "Sorted in Reverse Order (Bill #53 to #1, Serial No: 53 down to 1 — Most Recent First)",
    rec_desc,
    snos_desc
)

# --- TAB 2: CHRONOLOGICAL ORDER (Bill #1 to #53, Serial No: 1 to 53) ---
ws2 = wb.create_sheet(title="Serial No (1 to 53)")
rec_asc = sorted(records, key=lambda x: x['vch_idx'], reverse=False)
snos_asc = [r['vch_idx'] for r in rec_asc] # 1, 2, 3... 53
populate_sheet(
    ws2,
    "SKIL LIFESTYLE LLP — ALL 53 BILLS AUDIT REGISTER",
    "Sorted in Chronological Order (Bill #1 to #53, Serial No: 1 to 53)",
    rec_asc,
    snos_asc
)

# Save workbook to both project workspace and Downloads
output_proj = r'c:\Projects\SKIL Master\SKIL_53_Bills_Detailed_Audit.xlsx'
output_down = r'C:\Users\Sahil Kumar\Downloads\SKIL_53_Bills_Detailed_Audit.xlsx'

wb.save(output_proj)
wb.save(output_down)
print(f"Saved successfully to:\n1. {output_proj}\n2. {output_down}")
